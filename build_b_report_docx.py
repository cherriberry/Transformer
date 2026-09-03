from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / ".deps"))
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "B_LONGFORMER_MEMFORMER_EXPERIMENT_REPORT.docx"
DATA = ROOT / "results" / "b_role"

BLUE = "2E74B5"
DARK = "1F4D78"
INK = "0B2545"
MUTED = "5A6573"
FILL = "F2F4F7"
CALLOUT = "E8EEF5"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}"); tc_mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW"); tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_inches * 1440))); tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW"); tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360"); tbl_w.set(qn("w:type"), "dxa")
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd"); tbl_pr.append(ind)
    ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(round(w * 1440))); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(i, len(widths)-1)])
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header=True):
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs: set_run_font(run, size=9.5, color=INK)
            if ri == 0 and header:
                shade(cell, FILL)
                for p in cell.paragraphs:
                    for run in p.runs: run.bold = True; run.font.color.rgb = RGBColor.from_string(DARK)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)
    set_table_geometry(table, widths); style_table(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.1
    if align is not None: p.alignment = align
    r = p.add_run(text); set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); set_run_font(r, size={1:16,2:13,3:12}[level], color={1:BLUE,2:BLUE,3:DARK}[level], bold=True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text); set_run_font(r, size=10.5, color=INK)
    return p


def add_picture(doc, path, width=6.35):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(width))


def build():
    eff = json.loads((DATA / "efficiency.json").read_text(encoding="utf-8"))
    quality = json.loads((DATA / "quality_pilot.json").read_text(encoding="utf-8"))
    tiny = json.loads((DATA / "tinylm_quality.json").read_text(encoding="utf-8"))
    corr = json.loads((DATA / "correctness.json").read_text(encoding="utf-8"))
    abl = json.loads((DATA / "ablations.json").read_text(encoding="utf-8"))

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    # Base styles: standard_business_brief preset.
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for name, size, color, before, after in (("Heading 1",16,BLUE,16,8),("Heading 2",13,BLUE,12,6),("Heading 3",12,DARK,8,4)):
        st=doc.styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.line_spacing=1.1

    # Quiet running header/footer.
    hp = sec.header.paragraphs[0]; hp.text = "B 角色实验总结 | Longformer + Memformer"; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs: set_run_font(r, size=9, color=MUTED)
    fp = sec.footer.paragraphs[0]; fp.text = "高效 Transformer 比较研究 · 2026-09-02"; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs: set_run_font(r, size=9, color=MUTED)

    add_para(doc, "B 角色实验总结报告", size=23, color=INK, bold=True, after=4)
    add_para(doc, "Longformer 与 Memformer：架构、长程信息通路与效率效果调研", size=14, color="373737", after=14)
    meta = add_table(doc, ["项目", "内容"], [
        ("角色", "B：Longformer、Memformer 长程结构与记忆"),
        ("日期", "2026-09-02"),
        ("统一口径", "6 layers · hidden=384 · heads=8 · FFN=1536 · causal · BF16"),
        ("证据标签", "paper_reported / paper_aligned / matched_tinylm / innovation_validation"),
        ("执行状态", "M0–M4 完成；原始 JSON、图表和代码可追溯"),
    ], [1.35, 5.15])
    add_para(doc, "结论先行", size=11, color=DARK, bold=True, after=3)
    call = doc.add_table(rows=1, cols=1); set_table_geometry(call, [6.5]); shade(call.cell(0,0), CALLOUT); cell_margins(call.cell(0,0), top=120, bottom=120)
    p=call.cell(0,0).paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run("在本机 BF16、batch=1、N≤4096 的 attention-level 约束下，Longformer 在 N=4096 显现窗口化显存优势；Memformer 以固定容量 recurrent memory 获得最低显存。两者的优势都依赖任务、长度和状态容量，不能给出脱离条件的总冠军。")
    set_run_font(r, size=10.5, color=INK)

    add_heading(doc, "1. 研究目标与证据边界", 1)
    add_para(doc, "本轮按 Word 中 B 角色交付范围，完成结构图、优势/缺点及可证伪条件、正确性测试、最小 pilot、统一质量/速度/显存、关键消融、Passkey/Copy/Needle 机制任务和失败案例。论文公开设置与统一 TinyLM 结果分开报告：Longformer 原论文使用 text8/enwik8 character LM，分阶段扩展到 23,040；Memformer 原论文/官方 MemBART 使用 memory length 64/128、BART-base/large 主干。它们不是本项目 TinyLM，不能直接横比 PPL。", after=6)
    add_table(doc, ["方法", "主配置", "信息通路 / 复杂度", "本轮证据"], [
        ("Longformer", "window=128（odd width=129），global=1", "局部 O(NW) + 稀疏 global O(NG)", "matched_tinylm + mechanism"),
        ("Memformer", "segment=256，slots=64", "固定 M×d 状态；跨段融合与门控更新", "matched_tinylm + mechanism"),
    ], [1.0, 1.8, 2.35, 1.35])

    add_heading(doc, "2. 架构与数据流", 1)
    add_picture(doc, DATA / "architecture_dataflow.png", width=6.35)
    add_para(doc, "图 1  Longformer 的局部/全局路径与 Memformer 的跨段 memory 更新。", size=9.5, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_heading(doc, "2.1 Longformer", 2)
    for t in ["Q/K/V 投影后，每个 token 只在滑动窗口内计算 score；最大 score 张量为 [B,H,N,W]，不形成 N×N attention 矩阵。", "global token 作为稀疏远程汇聚点追加到普通 query 的候选集合；global query 另行读取所有有效 token，并避免局部/全局重复计数。", "causal=true 时窗口与 global 路径都施加上三角约束；padding query 输出置零。"]: add_bullet(doc,t)
    add_heading(doc, "2.2 Memformer", 2)
    for t in ["输入按 segment 切分；memory slots 与当前 segment 拼接后进行 fusion attention。", "memory rows 与 token rows 同时更新，memory 输出经 sigmoid gate 与旧状态融合，再传给下一段。", "state 元素数为 memory_slots×hidden_size，与总序列长度无关；reset 与 detach 是两个独立控制。"]: add_bullet(doc,t)

    add_heading(doc, "3. M1 正确性闸门", 1)
    add_para(doc, "原始记录：results/b_role/correctness.json。", size=9.5, color=MUTED, italic=True, after=4)
    add_table(doc, ["检查", "结果"], [
        ("Longformer 输出 shape=(2,37,32)", "通过；有限"),
        ("Longformer backward 梯度", "通过；有限"),
        ("full-window causal ≡ Standard", "MSE=0；relative L2=0；cosine=1"),
        ("Memformer 输出 shape=(1,16,32)", "通过；有限"),
        ("跨段 state delta vs reset", f"{corr['memformer_state_delta_vs_reset']:.5f}（非零，状态确实传播）"),
    ], [3.6, 2.9])
    add_para(doc, "仓库还保留 padding/boundary/global-key 去重/梯度/storage strategy 测试。当前环境未安装 pytest，因此本轮用同等 Python smoke assertions 复核关键闸门；测试代码位于 tests/test_longformer.py。", size=10.5, after=6)

    add_heading(doc, "4. M2 统一 TinyLM 质量 pilot", 1)
    add_para(doc, "同一 6 层/384 hidden/8 heads 主干，vocab=64、context=128、seed=17、12 个 AdamW steps，在确定性 copy-stream fallback 上训练和评估。该结果标记为 matched_tinylm，但数据集是 synthetic fallback；缺少 datasets/transformers 时未将 WikiText 伪装成正式结果。", after=6)
    tiny_rows=[]
    for row in tiny["records"]:
        tiny_rows.append((row["method"], f"{row['train_loss_last']:.4f}", f"{row['eval_nll']:.4f}", f"{row['eval_perplexity']:.4f}"))
    add_table(doc, ["方法", "train loss(last)", "eval NLL", "eval PPL"], tiny_rows, [1.5,1.65,1.45,1.9])
    add_para(doc, f"共享权重 attention proxy（N=256）补充结果：Longformer cosine={quality['longformer_vs_full_causal']['cosine_similarity']:.4f}、relative L2={quality['longformer_vs_full_causal']['relative_l2']:.4f}；Memformer cosine={quality['memformer_effect_vs_full_causal']['cosine_similarity']:.4f}、relative L2={quality['memformer_effect_vs_full_causal']['relative_l2']:.4f}。Memformer 改变信息通路，故按 effect 报告而非 approximation fidelity。", size=10.5, after=6)

    add_heading(doc, "5. M2 效率矩阵", 1)
    add_para(doc, "本机 CUDA / PyTorch 2.11.0 / BF16 / batch=1 / warmup=2 / timed runs=3 / seed=17。数值是 median latency 与 peak allocated。完整字段（reserved、原始 timed samples、status、config_hash）见 results/b_role/efficiency.json。", after=6)
    rows=[]
    by={}
    for r in eff["records"]:
        if r["status"]=="ok": by.setdefault(r["seq_len"],{})[r["method"]]=r
    for n in sorted(by):
        rr=by[n]; rows.append((n, f"{rr['Standard']['median_latency_s']*1000:.3f} / {rr['Standard']['peak_allocated_mb']:.1f}", f"{rr['Longformer']['median_latency_s']*1000:.3f} / {rr['Longformer']['peak_allocated_mb']:.1f}", f"{rr['Memformer']['median_latency_s']*1000:.3f} / {rr['Memformer']['peak_allocated_mb']:.1f}"))
    add_table(doc, ["N", "Standard ms / MiB", "Longformer ms / MiB", "Memformer ms / MiB"], rows, [0.65,1.85,2.0,2.0])
    add_picture(doc, DATA / "efficiency_curves.png", width=6.35)
    add_para(doc, "图 2  统一效率矩阵；短序列受 kernel/adapter 启动开销影响，N=4096 才明显显现窗口化和固定容量状态的优势。", size=9.5, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_bullet(doc, "N=4096：Longformer 10.758 ms / 78.6 MiB，相对 Standard 21.951 ms / 1333.5 MiB，约 2.04× 更快、94.1% 更省显存。")
    add_bullet(doc, "N=4096：Memformer 13.877 ms / 37.0 MiB，相对 Standard 约 1.58× 更快、97.2% 更省显存。")

    add_heading(doc, "6. M3 关键消融与机制任务", 1)
    add_heading(doc, "6.1 Longformer window / global", 2)
    add_para(doc, "N=2048、causal、seed=29。global=1 时 window 65→513，median latency 3.272→14.478 ms，peak allocated 约 141→860 MiB；global=0/1/4 的差异小于增大窗口的影响，但 global token 是远距汇聚路径的结构开关。", after=6)
    add_heading(doc, "6.2 Memformer segment / slots", 2)
    add_table(doc, ["segment", "slots=16 ms/MiB", "slots=64 ms/MiB", "slots=128 ms/MiB"], [
        (128,"24.722 / 30.0","22.019 / 30.7","22.287 / 33.5"),
        (256,"10.779 / 33.8","11.259 / 36.9","11.337 / 40.1"),
        (512,"5.740 / 50.8","5.345 / 55.1","4.964 / 61.4"),
    ], [0.85,1.85,1.9,1.9])
    add_picture(doc, DATA / "ablation_curves.png", width=6.35)
    add_para(doc, "图 3  关键结构旋钮消融；更长 segment 减少更新次数，更多 slots 提高状态容量但增加投影/显存成本。", size=9.5, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_heading(doc, "6.3 Passkey / Copy / Needle 与失败边界", 2)
    for t in ["Longformer local-only 在远距 probe 上不可达；增大 window 扩大局部传播范围；global token 提供稀疏汇聚但信息带宽有限。causal 版本只能向过去传播，不能泄漏未来。", "Memformer 始终有跨 segment state path，但容量固定：slots=16/64/128、hidden=384 时为 6,144/24,576/49,152 elements，即 BF16 状态约 12/48/96 KiB（每样本）。slots 太小会覆盖/遗忘；不 reset 会造成样本间污染；detach=true 会切断跨段反向梯度。", "这些机制记录属于 innovation_validation 的结构验证，不是训练后任务 accuracy。完整 reachability/state-capacity JSON 位于 results/b_role/mechanisms.json。"]: add_bullet(doc,t)

    add_heading(doc, "7. 假设检验、结论与下一步", 1)
    add_table(doc, ["方法", "优势假设 / 本轮证据", "缺点假设 / 证伪条件"], [
        ("Longformer", "O(NW+NG)；N=4096 latency/显存显著低于 full attention；window/global 可控", "固定窗口会漏远距依赖；global=0/1 在远距任务失败而增大 window/global 可恢复"),
        ("Memformer", "固定 M×d 状态；N=4096 显存近似 37 MiB；state delta 非零", "固定容量造成瓶颈/污染；slots 小或长段数下 Passkey/Needle 遗忘，reset/detach 改变结果"),
    ], [1.0,2.8,2.7])
    add_para(doc, "适用范围：Longformer 适合局部结构明显且需要少量全局汇聚的长序列；Memformer 适合可接受固定容量摘要、需要跨 segment 持久状态且显存预算严格的任务；短序列优先使用优化的 SDPA/full attention。", after=6)
    add_para(doc, "下一步：安装 datasets/transformers，使用 WikiText-103/PG-19 或原论文数据，扩展 seeds=17/29/43 与 50M–100M token 预算；Memformer 应使用完整 MemBART recurrent training，而不是 attention adapter。", after=8)

    add_heading(doc, "附录：可追溯文件与来源", 1)
    for t in ["实验计划：PLANS/B_LONGFORMER_MEMFORMER_EXPERIMENT_PLAN.md", "实验代码：b_role_experiments.py", "原始结果：results/b_role/{correctness.json,tinylm_quality.json,quality_pilot.json,efficiency.json,ablations.json,mechanisms.json}", "图表：results/b_role/{architecture_dataflow.png,efficiency_curves.png,ablation_curves.png}", "Longformer 原论文：https://arxiv.org/abs/2004.05150", "Memformer 原论文：https://arxiv.org/abs/2010.06891", "Memformer 官方实现：https://github.com/qywu/memformers", "参数标准：PAPER_CONFIGS_AND_RECOMMENDATIONS.md"]: add_bullet(doc,t)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
